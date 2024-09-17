from io import BytesIO
from pdfminer.high_level import extract_text
from pptx import Presentation

def extract_text_from_pptx(pptx_file):
    pptx_buffer = BytesIO()
    pptx_file.save(pptx_buffer)
    pptx_buffer.seek(0)
    prs = Presentation(pptx_buffer)
    extracted_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        extracted_text += run.text
    return extracted_text

def extract_text_from_pdf(pdf_file):
    pdf_buffer = BytesIO()
    pdf_file.save(pdf_buffer)
    pdf_buffer.seek(0)
    return extract_text(pdf_buffer)

def split_text_into_paragraphs(text, max_words=400):
    paragraphs = []
    words = text.split()
    while words:
        paragraphs.append(" ".join(words[:max_words]))
        words = words[max_words:]
    return paragraphs
